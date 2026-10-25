import io
import os
import re
import json
import requests
import openpyxl
from datetime import date, datetime
from urllib.parse import urlparse, parse_qs

# ==============================================================================
# PIPELINE DESIGN (Comment-Driven Development Blueprint)
# ==============================================================================
# phase 1: download excel sheet file into memory and read it
# phase 2: make a master list to store projects from all sheets
# Phase 3: loop through each sheet and 
    # 3a: skip sheets with "CLUSTER" or "LC X" in their name
    # 3b: scan rows with whitespace-normalized search to find header row
    # 3c: if there is no header row skip it
    # 3d: loop through data rows
        # - stop if reached "OUR EXPERIENCES" or "CURRENT STATE" section
        # - skip empty rows
        # - convert datetime/date cells to strings so JSON never crashes
        # - skip closed projects with "no slots now"
        # - clean google redirect links, fees, duration, and resolve SDGs
        # - keep valid projects (has title OR has link) and tag with country name
    # 3e: add the sheet's projects to the master list
# phase 4: save the data to frontend/src/data/projects.json using default=str
# ==============================================================================

# ID of the public spreadsheet
SHEET_ID = "10Fb0P3NT3sjRtycyCM-_85xWVa1oavWTfCVddh-6bwY"

# export Google Sheet as xlsx file
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

# target path for the React frontend
OUTPUT_PATH = "frontend/src/data/projects.json"

# map canonical project names to UN SDG numbers when sheets leave SDG blank
SDG_MAP = {
    "GREEN LEADERS": 13,
    "ROOTED": 15,
    "EQUIFY": 10,
    "FINGERPRINT": 4,
    "GLOBAL CLASSROOM": 4,
    "HEARTBEAT": 3,
    "SCALE UP": 17,
    "SCALE UP!": 17,
    "ON THE MAP": 8,
    "SKILL UP": 8,
    "SKILL UP!": 8,
    "AQUATICA": 14,
    "YOUTH 4 IMPACT": 4,
    "DISCOVER": 4,
    "RAISE YOUR VOICE": 5,
    "HAPPY BUS": 4,
    "EAT 4 CHANGE": 12,
    "EXPLORER": 15,
}


# ==============================================================================
# DECOUPLED SANITIZATION HELPERS
# ==============================================================================

def normalize_text(val):
    # normalizes newlines, non-breaking spaces (\xa0), tabs, and extra spaces
    if val is None:
        return ""
    return " ".join(str(val).replace("\xa0", " ").split()).strip().upper()


def find_field(record, *keywords):
    # flexible dictionary getter that matches column headers containing any of the keywords
    for key, val in record.items():
        if key and any(kw in key for kw in keywords):
            return val
    return None


def clean_url(raw_url):
    # unwrap google redirect links and reject plain text that is not a valid URL
    if not raw_url:
        return None
    url_str = str(raw_url).strip()
    if "google.com/url" in url_str:
        params = parse_qs(urlparse(url_str).query)
        url_str = params.get("q", [url_str])[0]

    match = re.search(r"(https?://[^\s)\"'>]+)", url_str)
    return match.group(1) if match else None


def clean_fee(raw_fee):
    # extract numeric float from messy strings like "75 EUR", "104,20€", "92.17 EUR", "190-450$"
    if raw_fee is None:
        return None
    match = re.search(r"(\d+(?:[.,]\d+)?)", str(raw_fee))
    if match:
        return float(match.group(1).replace(",", "."))
    return None


def clean_duration(raw_duration, raw_timeline=None):
    # extract primary integer weeks from "6 weeks", "6/4 weeks", or timeline fallback
    for source in (raw_duration, raw_timeline):
        if not source:
            continue
        match = re.search(r"(\d+)(?:\s*/\s*\d+)?\s*(?:week|w)", str(source), re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def clean_bool(raw_val):
    # standardize "Yes", "host families", "dorms of the NGO" -> True; "no", "not provided" -> False
    if not raw_val:
        return False
    val_str = str(raw_val).strip().upper()
    if val_str in ("", "NONE", "NO", "FALSE", "N/A", "-", ".") or val_str.startswith("NOT PROVIDED"):
        return False
    return True


def resolve_sdg(raw_sdg, project_name):
    # extract integer SDG from cell or fallback to SDG_MAP using project title
    if raw_sdg is not None:
        match = re.search(r"\d+", str(raw_sdg))
        if match:
            return int(match.group(0))
    cleaned_title = str(project_name or "").strip().upper()
    return SDG_MAP.get(cleaned_title)


# ==============================================================================
# MAIN ETL ORCHESTRATION
# ==============================================================================

def main():
    # --- PHASE 1: download excel sheet file into memory and read it ---
    print("Connecting to Google Sheets...")
    response = requests.get(EXPORT_URL, timeout=30)

    # check the response status code for errors
    if response.status_code != 200:
        print(f"Failed to download sheet. Status code: {response.status_code}")
        return

    print("Download complete. Parsing workbook in memory...")
    excel = io.BytesIO(response.content)
    wb = openpyxl.load_workbook(excel, data_only=True)

    # --- PHASE 2: make a master list to store projects from all sheets ---
    all_projects = []

    # --- PHASE 3: loop through each sheet ---
    for sheet in wb.worksheets:
        sheet_name = sheet.title.strip()

        # 3a: skip sheets with "CLUSTER" or "LC X" in their name
        if "CLUSTER" in sheet_name.upper() or "LC X" in sheet_name.upper():
            continue

        # 3b: scan rows using normalized whitespace to find header row
        header_row = None
        header_row_index = None
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row_normalized = [normalize_text(cell) for cell in row if cell is not None]

            # check if row contains both a link column and a project column
            has_link_col = any("LINK" in c for c in row_normalized)
            has_proj_col = any("PROJECT" in c for c in row_normalized)

            if has_link_col and has_proj_col:
                header_row = [normalize_text(col) if col is not None else None for col in row]
                header_row_index = row_index
                break

        # 3c: if there is no header row skip it
        if header_row is None:
            continue

        # 3d: loop through data rows
        projects = []
        for row_offset, row in enumerate(sheet.iter_rows(min_row=header_row_index + 1, values_only=True), start=1):
            # stop if reached "OUR EXPERIENCES" or "CURRENT STATE" section
            row_str = " ".join(normalize_text(c) for c in row if c is not None)
            if "OUR EXPERIENCES" in row_str or "CURRENT STATE" in row_str:
                break

            # skip empty rows
            if not any(row):
                continue

            # combine the header row to data and convert datetime objects to ISO strings
            project = {}
            for col_name, cell_value in zip(header_row, row):
                if col_name is not None:
                    if isinstance(cell_value, (date, datetime)):
                        cell_value = cell_value.isoformat()
                    project[col_name] = cell_value

            # skip closed projects marked with "no slots now"
            timeline_str = str(find_field(project, "TIMELINE") or "").lower()
            extra_info_str = str(find_field(project, "EXTRA INFO") or "").lower()
            if "no slots now" in timeline_str or "no slots now" in extra_info_str:
                continue

            # clean link and title using flexible header matching
            raw_link_val = find_field(project, "OPPORTUNITY LINK", "LINK")
            cleaned_link = clean_url(raw_link_val)

            raw_title_val = find_field(project, "PROJECT")
            raw_title = str(raw_title_val or "").strip()

            has_title = raw_title != "" and raw_title.upper() != "NONE"
            has_link = cleaned_link is not None

            # keep valid projects (either has a title OR has a valid link)
            if has_title or has_link:
                final_title = raw_title if has_title else "Global Volunteer Opportunity"

                raw_sdg_val = find_field(project, "SDG")
                raw_fee_val = find_field(project, "FEE", "COST")
                raw_duration_val = find_field(project, "DURATION")
                raw_booklet_val = find_field(project, "BOOKLET")
                raw_accomm_val = find_field(project, "ACCOMMOD", "ACCOMOD")

                # enrich and normalize fields for the React frontend contract
                project["ID"] = f"{sheet_name}_{header_row_index + row_offset}"
                project["COUNTRY"] = sheet_name
                project["PROJECT"] = final_title
                project["OPPORTUNITY LINK"] = cleaned_link
                project["BOOKLET"] = clean_url(raw_booklet_val)
                project["SDG"] = resolve_sdg(raw_sdg_val, final_title)
                project["FEE_NUMERIC"] = clean_fee(raw_fee_val)
                project["DURATION_WEEKS"] = clean_duration(raw_duration_val, timeline_str)
                project["ACCOMMODATION_BOOL"] = clean_bool(raw_accomm_val)

                projects.append(project)

        # 3e: add the sheet's projects to the master list
        all_projects.extend(projects)
        print(f"[{sheet_name}] Extracted {len(projects)} projects")

    # --- PHASE 4: save the data to a json file ---
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        # default=str guarantees datetime/custom objects serialize safely
        json.dump(all_projects, f, indent=2, ensure_ascii=False, default=str)

    print(f"\nDone! Saved {len(all_projects)} total projects to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()